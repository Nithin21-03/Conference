import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_template():
    os.makedirs('static/downloads', exist_ok=True)
    doc = docx.Document()

    # Set standard margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Header
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = 'International Conference on Big Data Tools and Techniques (ICBDTT-2026) | Sapthagiri NPS University'
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Conference title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run('Preparation of Papers for ICBDTT-2026: International Conference on Big Data Tools and Techniques')
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 16, 64) # SNPSU Navy

    # Authors block
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_auth = p_auth.add_run('First Author1, Second Author2, Corresponding Author*3\n')
    r_auth.font.name = 'Times New Roman'
    r_auth.font.size = Pt(11)
    r_auth.font.bold = True
    r_aff = p_auth.add_run('1,2,3 Department of Computer Science & Engineering, Sapthagiri NPS University, Bengaluru, India\n*Corresponding Author Email: author@university.edu\n')
    r_aff.font.name = 'Times New Roman'
    r_aff.font.size = Pt(10)
    r_aff.font.italic = True

    # Abstract Heading
    p_abs_hd = doc.add_paragraph()
    r_abs_hd = p_abs_hd.add_run('Abstract— ')
    r_abs_hd.font.bold = True
    r_abs_hd.font.name = 'Times New Roman'
    r_abs_txt = p_abs_hd.add_run('These instructions give you guidelines for preparing papers for the International Conference on Big Data Tools and Techniques (ICBDTT-2026) organized by Sapthagiri NPS University (SNPSU), Bengaluru. The abstract should be self-contained and between 150-250 words. It must reflect the context, objectives, methodology, key findings, and implications of your big data, artificial intelligence, cloud analytics, or data science research.')
    r_abs_txt.font.name = 'Times New Roman'

    # Keywords
    p_kw = doc.add_paragraph()
    r_kw_hd = p_kw.add_run('Keywords— ')
    r_kw_hd.font.bold = True
    r_kw_hd.font.name = 'Times New Roman'
    r_kw_txt = p_kw.add_run('Big Data, Distributed Computing, Machine Learning, Apache Spark, Cloud Analytics, Data Science, Neural Networks.')
    r_kw_txt.font.italic = True
    r_kw_txt.font.name = 'Times New Roman'

    # Headings
    doc.add_heading('I. INTRODUCTION', level=1)
    doc.add_paragraph('This document is a format template for Microsoft Word. Authors submitting to ICBDTT-2026 should strictly adhere to this guideline. Submitted manuscripts must be original research not previously published or under consideration elsewhere. The maximum length for full papers is 6 to 8 pages including all figures, tables, and bibliographic references.')

    doc.add_heading('II. RESEARCH METHODOLOGY & BIG DATA ARCHITECTURE', level=1)
    doc.add_paragraph('Describe the proposed system model, data ingestion pipeline, streaming or batch architecture (e.g., Apache Spark, Hadoop, Kafka, Ray), data lakes, machine learning or deep learning formulations, and privacy considerations.')

    doc.add_heading('III. EXPERIMENTAL EVALUATION & RESULTS', level=1)
    doc.add_paragraph('Provide empirical results, benchmarks, scalability assessments, latency and throughput measurements, and statistical comparisons against state-of-the-art baselines.')

    doc.add_heading('IV. CONCLUSION & FUTURE WORK', level=1)
    doc.add_paragraph('Summarize the primary research contributions and highlight potential future extensions.')

    doc.add_heading('REFERENCES', level=1)
    doc.add_paragraph('[1] J. Dean and S. Ghemawat, "MapReduce: Simplified Data Processing on Large Clusters," Communications of the ACM, vol. 51, no. 1, pp. 107-113, 2008.\n[2] M. Zaharia et al., "Apache Spark: A Unified Engine for Big Data Processing," Communications of the ACM, 2016.\n[3] Y. LeCun, Y. Bengio, and G. Hinton, "Deep learning," Nature, vol. 521, no. 7553, pp. 436-444, 2015.')

    filepath = os.path.join('static', 'downloads', 'SNPSU_BDTT_2026_Paper_Template.docx')
    doc.save(filepath)
    print(f"Generated official template at: {filepath}, size: {os.path.getsize(filepath)} bytes")

if __name__ == '__main__':
    create_template()
